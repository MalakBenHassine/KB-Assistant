"""A small corpus in the four supported formats, generated in code.

Generating instead of committing binaries keeps the fixtures reviewable: the
content of every test document is visible in this file.
"""

import io
import zipfile

import docx

LEAVE_POLICY_MD = """\
---
owner: hr
---
# Leave policy

Every employee is entitled to 25 days of paid annual leave per calendar year.
Unused days can be carried over until March 31 of the following year.

## Parental leave

### Eligibility

Employees with at least six months of service are eligible for 16 weeks of
fully paid parental leave.

### How to apply

1. Tell your manager at least eight weeks in advance.
2. Submit the request in the HR portal.

| Leave type | Days |
| --- | --- |
| Annual | 25 |
| Sick (self-certified) | 5 |

```
portal: https://hr.example/leave
```
"""

ONCALL_RUNBOOK_HTML = """\
<!doctype html>
<html>
<head><title>On-call runbook</title><script>steal()</script><style>p{}</style></head>
<body>
<nav><a href="/">Home</a> | <a href="/runbooks">Runbooks</a></nav>
<h1>Database failover</h1>
<p>Page the database on-call engineer before starting a failover.</p>
<h2>Steps</h2>
<ol>
  <li><p>Freeze deployments in the release channel.</p></li>
  <li>Promote the replica with <code>pg_ctl promote</code>.</li>
</ol>
<pre>kubectl -n data rollout restart deploy/api</pre>
<img src="https://attacker.example/pixel.png" alt="">
<footer>Engineering handbook</footer>
</body>
</html>
"""


def make_docx() -> bytes:
    document = docx.Document()
    document.add_heading("Expense policy", level=1)
    document.add_paragraph("Expenses above 500 EUR need written approval from a director.")
    document.add_heading("Travel", level=2)
    document.add_paragraph("Book trains for journeys under four hours.", style="List Bullet")
    document.add_paragraph("Economy class only on flights under six hours.", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Meal"
    table.cell(0, 1).text = "40 EUR per day"
    table.cell(1, 0).text = "Hotel"
    table.cell(1, 1).text = "150 EUR per night"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(pages: list[list[str]]) -> bytes:
    """A minimal valid PDF: one text line per entry, Helvetica, one page per list."""
    objects: list[bytes] = []
    page_ids = [4 + 2 * i for i in range(len(pages))]
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, lines in enumerate(pages):
        content_id = page_ids[i] + 1
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>".encode()
        )
        ops = "".join(
            f"BT /F1 11 Tf 72 {740 - 16 * n} Td ({_pdf_escape(line)}) Tj ET\n"
            for n, line in enumerate(lines)
        ).encode()
        objects.append(b"<< /Length %d >>\nstream\n" % len(ops) + ops + b"endstream")

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    )
    return out.getvalue()


def handbook_pdf() -> bytes:
    """Three pages with a running header and footer, which cleaning must remove."""
    return make_pdf(
        [
            [
                "Acme Employee Handbook",
                "Welcome to Acme.",
                "Our values are trust and craft.",
                "Page 1 of 3",
            ],
            [
                "Acme Employee Handbook",
                "Working hours are flexible between 8 and 18.",
                "Core hours are 10 to 16.",
                "Page 2 of 3",
            ],
            [
                "Acme Employee Handbook",
                "Security training is mandatory every year.",
                "Report incidents to security@acme.example.",
                "Page 3 of 3",
            ],
        ]
    )


def zip_bomb_docx() -> bytes:
    """A 'DOCX' whose main part decompresses to 20 MiB of zeros."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", b"\0" * (20 * 1024 * 1024))
    return buffer.getvalue()
